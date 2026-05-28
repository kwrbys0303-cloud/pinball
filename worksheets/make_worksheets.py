"""
アセスメントプリント生成スクリプト
特別支援学校高等部向け（現状把握・配置確認）
"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

# ========== ユーティリティ ==========

def set_cell_border(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for border_name in ["top","left","bottom","right"]:
        border = OxmlElement(f"w:{border_name}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "6")
        border.set(qn("w:color"), "000000")
        tcPr.append(border)

def header(doc, title, subtitle=""):
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(20)

    if subtitle:
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run(subtitle)
        r2.font.size = Pt(11)

    # 名前・日付欄
    tbl = doc.add_table(rows=1, cols=4)
    tbl.style = "Table Grid"
    labels = ["なまえ", "　　　　　　　　　　", "日づけ", "　　　　　　　"]
    for i, (cell, label) in enumerate(zip(tbl.rows[0].cells, labels)):
        cell.text = label
        cell.paragraphs[0].runs[0].font.size = Pt(11)
        set_cell_border(cell)
    doc.add_paragraph()

def section(doc, num, text):
    p = doc.add_paragraph()
    run = p.add_run(f"【{num}】 {text}")
    run.bold = True
    run.font.size = Pt(13)

def blank(doc, n=1):
    for _ in range(n):
        doc.add_paragraph()

def answer_line(doc, label, width="　" * 10):
    p = doc.add_paragraph()
    p.add_run(f"　{label}　　{width}")
    p.runs[-1].underline = True

# ========== 1. 計算プリント ==========

def make_keisan():
    doc = Document()
    for section_obj in doc.sections:
        section_obj.top_margin = Cm(1.5)
        section_obj.bottom_margin = Cm(1.5)
        section_obj.left_margin = Cm(2)
        section_obj.right_margin = Cm(2)

    header(doc, "けいさんプリント", "（計算のアセスメント）")

    section(doc, 1, "たしざん　つぎのけいさんをしましょう。")
    problems = [
        ("３ ＋ ５ ＝", "２ ＋ ７ ＝", "４ ＋ ４ ＝"),
        ("１６ ＋ ５ ＝", "２４ ＋ ８ ＝", "３７ ＋ ６ ＝"),
    ]
    for row in problems:
        tbl = doc.add_table(rows=1, cols=3)
        tbl.style = "Table Grid"
        for cell, prob in zip(tbl.rows[0].cells, row):
            cell.text = prob
            cell.paragraphs[0].runs[0].font.size = Pt(14)
            cell.paragraphs[0].paragraph_format.space_after = Pt(8)
            set_cell_border(cell)
        doc.add_paragraph()

    section(doc, 2, "ひきざん　つぎのけいさんをしましょう。")
    problems2 = [
        ("９ － ３ ＝", "８ － ５ ＝", "７ － ７ ＝"),
        ("３２ － ７ ＝", "４１ － ８ ＝", "５０ － ６ ＝"),
    ]
    for row in problems2:
        tbl = doc.add_table(rows=1, cols=3)
        tbl.style = "Table Grid"
        for cell, prob in zip(tbl.rows[0].cells, row):
            cell.text = prob
            cell.paragraphs[0].runs[0].font.size = Pt(14)
            set_cell_border(cell)
        doc.add_paragraph()

    section(doc, 3, "かけざん　つぎのけいさんをしましょう。")
    problems3 = [
        ("２ × ３ ＝", "４ × ２ ＝", "３ × ３ ＝"),
        ("５ × ４ ＝", "６ × ３ ＝", "２ × ９ ＝"),
    ]
    for row in problems3:
        tbl = doc.add_table(rows=1, cols=3)
        tbl.style = "Table Grid"
        for cell, prob in zip(tbl.rows[0].cells, row):
            cell.text = prob
            cell.paragraphs[0].runs[0].font.size = Pt(14)
            set_cell_border(cell)
        doc.add_paragraph()

    section(doc, 4, "もんだい　よんでこたえましょう。")
    doc.add_paragraph()
    q1 = doc.add_paragraph()
    q1.add_run("① みかんが ６こ あります。３こ たべました。のこりは なんこですか？").font.size = Pt(13)
    doc.add_paragraph()
    ans1 = doc.add_paragraph()
    ans1.add_run("　　しき：　　　　　　　　　　　　こたえ：　　　　　こ").font.size = Pt(13)
    doc.add_paragraph()
    q2 = doc.add_paragraph()
    q2.add_run("② えんぴつが １ダース（１２ほん）あります。１クラスに ３ほんずつ くばります。").font.size = Pt(13)
    q2b = doc.add_paragraph()
    q2b.add_run("　　なんクラスに くばれますか？").font.size = Pt(13)
    doc.add_paragraph()
    ans2 = doc.add_paragraph()
    ans2.add_run("　　しき：　　　　　　　　　　　　こたえ：　　　　　クラス").font.size = Pt(13)

    doc.save("/home/user/pinball/worksheets/01_けいさんプリント.docx")
    print("01 done")

# ========== 2. 図形プリント ==========

def make_zukei():
    doc = Document()
    for section_obj in doc.sections:
        section_obj.top_margin = Cm(1.5)
        section_obj.bottom_margin = Cm(1.5)
        section_obj.left_margin = Cm(2)
        section_obj.right_margin = Cm(2)

    header(doc, "ずけいプリント", "（図形のアセスメント）")

    section(doc, 1, "なまえをかきましょう。　下の図形の名前を書きましょう。")
    doc.add_paragraph()

    # 図形の説明テキストで代替（図なし）
    shapes_q = [
        ("①", "かどが 3つ、へんが 3本の形", "　　　　　　　　（　　　　　　　　　　）"),
        ("②", "かどが 4つ、へんが 4本の形（よこが ながい）", "　（　　　　　　　　　　）"),
        ("③", "かどが 4つ、全部のへんが同じ長さの形", "　（　　　　　　　　　　）"),
        ("④", "まるい形（かどが ない）", "　（　　　　　　　　　　）"),
        ("⑤", "かどが 5つ、へんが 5本の形", "　（　　　　　　　　　　）"),
    ]
    tbl = doc.add_table(rows=len(shapes_q), cols=3)
    tbl.style = "Table Grid"
    widths = [Cm(1), Cm(9), Cm(5)]
    for i, (num, desc, ans) in enumerate(shapes_q):
        row = tbl.rows[i]
        row.cells[0].text = num
        row.cells[1].text = desc
        row.cells[2].text = ans
        for j, cell in enumerate(row.cells):
            cell.paragraphs[0].runs[0].font.size = Pt(12)
            set_cell_border(cell)
    doc.add_paragraph()

    section(doc, 2, "かきましょう。　下のわくに、言われた形をかきましょう。")
    doc.add_paragraph()

    shapes_draw = ["さんかくけい（三角形）", "しかくけい（四角形）", "ごかくけい（五角形）", "まる（円）"]
    tbl2 = doc.add_table(rows=1, cols=4)
    tbl2.style = "Table Grid"
    for cell, label in zip(tbl2.rows[0].cells, shapes_draw):
        p = cell.paragraphs[0]
        p.add_run(label).font.size = Pt(11)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        # 空白行を追加して描画スペース確保
        for _ in range(6):
            cell.add_paragraph()
        set_cell_border(cell)
    doc.add_paragraph()

    section(doc, 3, "えらびましょう。　正しいものを　○　でかこみましょう。")
    doc.add_paragraph()
    qs = [
        "① 三角形のかどの数は　　（ ２ ・ ３ ・ ４ ）つ　です。",
        "② 四角形のへんの数は　　（ ３ ・ ４ ・ ５ ）本　です。",
        "③ 円には かどが　　（ ある ・ ない ）　。",
    ]
    for q in qs:
        p = doc.add_paragraph(q)
        p.runs[0].font.size = Pt(13)
        doc.add_paragraph()

    doc.save("/home/user/pinball/worksheets/02_ずけいプリント.docx")
    print("02 done")

# ========== 3. 時間プリント ==========

def make_jikan():
    doc = Document()
    for section_obj in doc.sections:
        section_obj.top_margin = Cm(1.5)
        section_obj.bottom_margin = Cm(1.5)
        section_obj.left_margin = Cm(2)
        section_obj.right_margin = Cm(2)

    header(doc, "じかんプリント", "（時間のアセスメント）")

    section(doc, 1, "とけいをよみましょう。　何時何分ですか。")
    doc.add_paragraph()

    clock_descs = [
        ("①", "みじかいはりが「３」、ながいはりが「１２」をさしています。", "（　　　　じ　　　　ふん）"),
        ("②", "みじかいはりが「７」と「８」のあいだ、ながいはりが「６」をさしています。", "（　　　　じ　　　　ふん）"),
        ("③", "みじかいはりが「１１」にちかい、ながいはりが「９」をさしています。", "（　　　　じ　　　　ふん）"),
        ("④", "みじかいはりが「２」と「３」のあいだ、ながいはりが「３」をさしています。", "（　　　　じ　　　　ふん）"),
    ]
    tbl = doc.add_table(rows=len(clock_descs), cols=3)
    tbl.style = "Table Grid"
    for i, (num, desc, ans) in enumerate(clock_descs):
        row = tbl.rows[i]
        row.cells[0].text = num
        row.cells[1].text = desc
        row.cells[2].text = ans
        for cell in row.cells:
            cell.paragraphs[0].runs[0].font.size = Pt(12)
            set_cell_border(cell)
    doc.add_paragraph()

    section(doc, 2, "あいだの時間　何分かかりましたか。")
    doc.add_paragraph()
    time_qs = [
        "① 午前 ９時００分 に 家を でました。\n　　学校に ９時３０分 につきました。\n　　何分 かかりましたか？　　　　（　　　　ふん）",
        "② 午後 ２時００分 から そうじが はじまりました。\n　　２時２０分 におわりました。\n　　そうじは 何分でしたか？　　　（　　　　ふん）",
        "③ えいがが 午後 １時１５分 にはじまり、\n　　３時００分 におわりました。\n　　えいがは 何時間何分でしたか？　（　　　　じかん　　　　ふん）",
    ]
    for q in time_qs:
        p = doc.add_paragraph(q)
        p.runs[0].font.size = Pt(12)
        doc.add_paragraph()

    section(doc, 3, "ならべましょう。　時刻を早い順に数字でならべましょう。")
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.add_run("　 ア　午後３時　　　イ　午前９時　　　ウ　正午（12時）\n\n　 早い順：（　　　）→（　　　）→（　　　）").font.size = Pt(13)

    doc.save("/home/user/pinball/worksheets/03_じかんプリント.docx")
    print("03 done")

# ========== 4. 言葉プリント ==========

def make_kotoba():
    doc = Document()
    for section_obj in doc.sections:
        section_obj.top_margin = Cm(1.5)
        section_obj.bottom_margin = Cm(1.5)
        section_obj.left_margin = Cm(2)
        section_obj.right_margin = Cm(2)

    header(doc, "ことばプリント", "（言葉のアセスメント）")

    section(doc, 1, "いみをえらびましょう。　言葉の意味として正しいものを選びましょう。")
    doc.add_paragraph()
    vocab_qs = [
        ("① 「やさしい」", [
            "ア　むずかしい　　イ　おだやかで おもいやりがある　　ウ　うるさい"
        ], "（　　　）"),
        ("② 「しずか」", [
            "ア　おとが あまりしない　　イ　おおきな こえ　　ウ　はやい"
        ], "（　　　）"),
        ("③ 「うれしい」", [
            "ア　いたい　　イ　こわい　　ウ　たのしい きもち"
        ], "（　　　）"),
    ]
    for word, choices, ans in vocab_qs:
        p1 = doc.add_paragraph(word)
        p1.runs[0].font.size = Pt(13)
        p1.runs[0].bold = True
        p2 = doc.add_paragraph("　　" + choices[0] + "　　" + ans)
        p2.runs[0].font.size = Pt(12)
        doc.add_paragraph()

    section(doc, 2, "はんたいのことば　（　）に反対の言葉を書きましょう。")
    doc.add_paragraph()
    opposites = [
        ("① 大きい", "（　　　　　　）"),
        ("② 新しい", "（　　　　　　）"),
        ("③ 上", "（　　　　　　）"),
        ("④ 入る", "（　　　　　　）"),
    ]
    tbl = doc.add_table(rows=2, cols=4)
    tbl.style = "Table Grid"
    for i, (word, ans) in enumerate(opposites):
        col = i % 4
        row_idx = i // 4
        cell = tbl.rows[row_idx].cells[col]
        cell.text = f"{word} → {ans}"
        cell.paragraphs[0].runs[0].font.size = Pt(12)
        set_cell_border(cell)
    doc.add_paragraph()

    section(doc, 3, "ことばをあてはめましょう。　（　）に合う言葉を下から選びましょう。")
    doc.add_paragraph()
    sentences = [
        "① 今日は 天気が いいので、そとで （　　　　　）ましょう。",
        "② ごはんを たべる まえに （　　　　　）を あらいます。",
        "③ 本を （　　　　　）が おわったら、たなに もどします。",
    ]
    for s in sentences:
        p = doc.add_paragraph(s)
        p.runs[0].font.size = Pt(13)
        doc.add_paragraph()

    p_choices = doc.add_paragraph()
    p_choices.add_run("　　【あそび　　て　　よみ】").font.size = Pt(13)
    p_choices.runs[0].bold = True

    doc.save("/home/user/pinball/worksheets/04_ことばプリント.docx")
    print("04 done")

# ========== 5. 漢字プリント ==========

def make_kanji():
    doc = Document()
    for section_obj in doc.sections:
        section_obj.top_margin = Cm(1.5)
        section_obj.bottom_margin = Cm(1.5)
        section_obj.left_margin = Cm(2)
        section_obj.right_margin = Cm(2)

    header(doc, "かんじプリント", "（漢字のアセスメント）")

    section(doc, 1, "よみましょう。　漢字の読み方をひらがなで書きましょう。")
    doc.add_paragraph()

    read_qs = [
        ("① 山", "（　　　　）"),
        ("② 川", "（　　　　）"),
        ("③ 学校", "（　　　　）"),
        ("④ 時間", "（　　　　）"),
        ("⑤ 図書館", "（　　　　）"),
        ("⑥ 勉強", "（　　　　）"),
    ]
    tbl = doc.add_table(rows=3, cols=4)
    tbl.style = "Table Grid"
    all_qs = read_qs
    for i, (kanji, ans) in enumerate(all_qs):
        r, c = divmod(i, 4)
        if r < 3 and c < 4:
            cell = tbl.rows[r].cells[c]
            cell.text = f"{kanji}　{ans}"
            cell.paragraphs[0].runs[0].font.size = Pt(14)
            set_cell_border(cell)
    doc.add_paragraph()

    section(doc, 2, "かきましょう。　読み方に合う漢字を書きましょう。")
    doc.add_paragraph()

    write_qs = [
        ("① やま", "（　　）"),
        ("② かわ", "（　　）"),
        ("③ て", "（　　）"),
        ("④ め", "（　　）"),
        ("⑤ ひ（太陽）", "（　　）"),
        ("⑥ き（木）", "（　　）"),
        ("⑦ いぬ", "（　　）"),
        ("⑧ はな（花）", "（　　）"),
    ]
    tbl2 = doc.add_table(rows=4, cols=4)
    tbl2.style = "Table Grid"
    for i, (yomi, ans) in enumerate(write_qs):
        r, c = divmod(i, 4)
        cell = tbl2.rows[r].cells[c]
        cell.text = f"{yomi}　{ans}"
        cell.paragraphs[0].runs[0].font.size = Pt(13)
        set_cell_border(cell)
    doc.add_paragraph()

    section(doc, 3, "ぶんをよみましょう。　下線の漢字の読み方を書きましょう。")
    doc.add_paragraph()
    sentences = [
        "① 毎朝、学校に 行きます。　　　　（まい　　　）（　　　こう）",
        "② 友達と 公園で 遊びました。　　（　　　だち）（こう　　　）",
        "③ 先生に 質問しました。　　　　　（　　　せい）（　　　もん）",
    ]
    for s in sentences:
        p = doc.add_paragraph(s)
        p.runs[0].font.size = Pt(12)
        doc.add_paragraph()

    doc.save("/home/user/pinball/worksheets/05_かんじプリント.docx")
    print("05 done")

# ========== 6. 言葉遣いプリント ==========

def make_kotobazukai():
    doc = Document()
    for section_obj in doc.sections:
        section_obj.top_margin = Cm(1.5)
        section_obj.bottom_margin = Cm(1.5)
        section_obj.left_margin = Cm(2)
        section_obj.right_margin = Cm(2)

    header(doc, "ことばづかいプリント", "（言葉遣いのアセスメント）")

    section(doc, 1, "ていねいに言いましょう。　（　）にていねいな言い方を書きましょう。")
    doc.add_paragraph()
    polite_qs = [
        ("① 「食べる」", "→ ていねいに：（　　　　　　　　　　　　　　　）"),
        ("② 「行く」", "→ ていねいに：（　　　　　　　　　　　　　　　）"),
        ("③ 「わからない」", "→ ていねいに：（　　　　　　　　　　　　　　　）"),
        ("④ 「ちょっと待って」", "→ ていねいに：（　　　　　　　　　　　　　　　）"),
    ]
    for word, ans in polite_qs:
        tbl = doc.add_table(rows=1, cols=2)
        tbl.style = "Table Grid"
        tbl.rows[0].cells[0].text = word
        tbl.rows[0].cells[0].paragraphs[0].runs[0].font.size = Pt(13)
        tbl.rows[0].cells[1].text = ans
        tbl.rows[0].cells[1].paragraphs[0].runs[0].font.size = Pt(13)
        for cell in tbl.rows[0].cells:
            set_cell_border(cell)
        doc.add_paragraph()

    section(doc, 2, "だれに・どんな話し方？　線でむすびましょう。")
    doc.add_paragraph()
    p_inst = doc.add_paragraph()
    p_inst.add_run("　左の「話す相手」と、右の「言い方」を線でむすびましょう。").font.size = Pt(12)
    doc.add_paragraph()

    tbl2 = doc.add_table(rows=4, cols=3)
    tbl2.style = "Table Grid"
    left_col = ["① 先生に", "② 友だちに", "③ お店の人に", "④ 家族（おとうさん・おかあさん）に"]
    mid_col = ["", "", "", ""]
    right_col = [
        "ア　「ねえねえ、いっしょに行こうよ」",
        "イ　「すみません、これをください」",
        "ウ　「ただいまー！ごはんまだ？」",
        "エ　「先生、質問があります」",
    ]
    for i in range(4):
        tbl2.rows[i].cells[0].text = left_col[i]
        tbl2.rows[i].cells[0].paragraphs[0].runs[0].font.size = Pt(12)
        tbl2.rows[i].cells[1].text = "（　→　　）"
        tbl2.rows[i].cells[1].paragraphs[0].runs[0].font.size = Pt(12)
        tbl2.rows[i].cells[2].text = right_col[i]
        tbl2.rows[i].cells[2].paragraphs[0].runs[0].font.size = Pt(12)
        for cell in tbl2.rows[i].cells:
            set_cell_border(cell)
    doc.add_paragraph()

    section(doc, 3, "どちらがよい？　場面に合う言い方を選び、○をつけましょう。")
    doc.add_paragraph()
    scene_qs = [
        ("① 授業中、先生が話しているときに トイレに行きたくなりました。",
         "ア　だまって席を立つ。\nイ　「先生、トイレに行っていいですか」と言う。"),
        ("② お店で買い物をして、お金をわたすとき。",
         "ア　「はい」と言ってわたす。\nイ　だまってカウンターにおく。"),
    ]
    for question, choices in scene_qs:
        p = doc.add_paragraph(question)
        p.runs[0].font.size = Pt(12)
        p.runs[0].bold = True
        p2 = doc.add_paragraph(choices)
        p2.runs[0].font.size = Pt(12)
        doc.add_paragraph()

    doc.save("/home/user/pinball/worksheets/06_ことばづかいプリント.docx")
    print("06 done")

# ========== 実行 ==========

make_keisan()
make_zukei()
make_jikan()
make_kotoba()
make_kanji()
make_kotobazukai()

print("\n全プリント生成完了！")
